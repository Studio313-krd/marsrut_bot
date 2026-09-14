import asyncio
import io
import zipfile
from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock

import httpx
import pytest
from fastapi import FastAPI
from openpyxl import load_workbook

from app.leadership import router


@pytest.fixture
def integration():
    app = FastAPI()
    app.include_router(router)
    app.state.settings = SimpleNamespace(leadership_report_key="k" * 40)
    payload = {
        "schemaVersion": 1,
        "items": [],
        "total": 0,
        "webActors": [],
        "generatedAt": "2026-09-14T10:00:00Z",
        "period": {"from": "2026-09-01T00:00:00Z", "to": "2026-09-14T10:00:00Z"},
    }
    service = SimpleNamespace(_statistics_lock=asyncio.Lock(), storage=Mock(), site=Mock())
    service.storage.list_admins.return_value = []
    service.site.request_statistics = AsyncMock(return_value=payload)
    app.state.service = service
    return app, service


@pytest.mark.asyncio
async def test_leadership_auth_range_and_real_workbook(integration):
    app, service = integration
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as client:
        assert (await client.get("/api/leadership/v1/catalog")).status_code == 401
        client.headers["Authorization"] = "Bearer " + "k" * 40
        assert (await client.get("/api/leadership/v1/catalog")).json()["project"] == "marshrut"
        body = {
            "fromUtc": "2026-09-01T00:00:00Z",
            "toUtc": "2026-09-14T10:00:00Z",
            "timeZone": "Europe/Moscow",
            "filters": {},
        }
        response = await client.post("/api/leadership/v1/report", json=body)
        assert response.status_code == 200
        with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
            workbook = load_workbook(io.BytesIO(archive.read("report.xlsx")), data_only=True)
            assert len(workbook.sheetnames) == 8
            assert "События" in workbook.sheetnames
            assert "updatedAt" in archive.read("prompt.txt").decode()
        service.site.request_statistics.assert_awaited_once_with(
            from_date="2026-09-01T00:00:00+00:00", to_date="2026-09-14T10:00:00+00:00"
        )
        assert not service._statistics_lock.locked()
        assert (
            await client.post("/api/leadership/v1/report", json={**body, "filters": {"x": "y"}})
        ).status_code == 422
        assert (
            await client.post("/api/leadership/v1/report", json={**body, "fromUtc": "2026-09-01"})
        ).status_code == 422
        async with service._statistics_lock:
            assert (await client.post("/api/leadership/v1/report", json=body)).status_code == 429
        app.state.settings.leadership_report_key = ""
        assert (await client.post("/api/leadership/v1/report", json=body)).status_code == 404
