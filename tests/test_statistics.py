from __future__ import annotations

import io
from datetime import UTC, datetime, timedelta
from zoneinfo import ZoneInfo

import openpyxl
import pytest

from app.statistics import build_statistics_workbook, request_facts


def activity(
    number,
    minutes,
    kind,
    *,
    actor_type="ADMIN",
    channel="TELEGRAM",
    actor="bot-admin",
    name="Борис",
    status=None,
    previous=None,
):
    return {
        "id": f"activity-{number}",
        "createdAt": (datetime(2026, 9, 10, 21, 30, tzinfo=UTC) + timedelta(minutes=minutes)).isoformat(),
        "type": kind,
        "actorKey": actor,
        "actorName": name,
        "actorType": actor_type,
        "actorChannel": channel,
        "fromStatus": previous,
        "toStatus": status,
    }


def sample_payload():
    item = {
        "id": "request-1",
        "requestNumber": "MP-260910-TEST01",
        "createdAt": "2026-09-10T21:30:00Z",
        "updatedAt": "2026-09-11T02:30:00Z",
        "source": "WEBSITE",
        "status": "COMPLETED",
        "company": '=HYPERLINK("https://example.test","x")',
        "assignedAdminKey": "someone-else",
        "assignedAdminName": "Другой ответственный",
        "activities": [
            activity(0, 0, "CREATED", actor_type="USER", channel="WEBSITE"),
            activity(1, 5, "MESSAGE_FROM_USER", actor_type="USER"),
            activity(2, 7, "COMMENT", actor_type="SYSTEM", channel="SYSTEM", actor=None),
            activity(3, 9, "COMMENT", actor_type="USER"),
            activity(4, 60, "ASSIGNED"),
            activity(5, 90, "STATUS_CHANGED", channel="MAX", previous="NEW", status="IN_PROGRESS"),
            activity(
                6,
                180,
                "STATUS_CHANGED",
                channel="WEBSITE",
                actor="web-admin",
                name="Анна",
                previous="IN_PROGRESS",
                status="COMPLETED",
            ),
            activity(7, 240, "STATUS_CHANGED", previous="COMPLETED", status="IN_PROGRESS"),
            activity(8, 300, "STATUS_CHANGED", previous="IN_PROGRESS", status="COMPLETED"),
        ],
    }
    legacy = {
        "id": "request-2",
        "requestNumber": "MP-OLD",
        "createdAt": "2026-09-12T00:00:00Z",
        "updatedAt": "2026-09-14T10:00:00Z",
        "source": "MAX",
        "status": "COMPLETED",
        "activities": [],
    }
    return {
        "schemaVersion": 1,
        "generatedAt": "2026-09-14T12:00:00Z",
        "period": {"from": None, "to": "2026-09-14T12:00:00Z"},
        "total": 2,
        "items": [legacy, item],
        "webActors": [{"id": "web-admin", "name": "Анна"}],
    }


def test_metrics_use_real_actors_and_events_not_assignee_or_updated_at():
    payload = sample_payload()
    content = build_statistics_workbook(
        payload, [{"id": "bot-admin", "display_name": "Борис"}], ZoneInfo("Europe/Moscow")
    )
    values = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    formulas = openpyxl.load_workbook(io.BytesIO(content), data_only=False)
    ws = values["Заявки"]
    assert ws["M2"].value == "Борис"
    assert ws["N2"].value == "TELEGRAM"
    assert ws["O2"].value == 60
    assert ws["S2"].value == "Анна"
    assert ws["T2"].value == "WEBSITE"
    assert ws["U2"].value == 180
    assert ws["V2"].value == 90
    assert ws["Y2"].value == "Борис"
    assert ws["AC2"].value == 5
    assert ws["D2"].value == datetime(2026, 9, 11, 0, 30)
    assert ws["O3"].value in (None, "")
    assert ws["U3"].value in (None, "")
    assert ws["AB3"].value == 0
    assert formulas["Заявки"]["O2"].data_type == "f"
    assert formulas["Заявки"]["G2"].data_type == "s"
    assert values["Сводка"]["B2"].value == 2
    assert values["Сводка"]["B3"].value == 4
    assert values["Сводка"]["B4"].value == 0.5
    days = {row[0].date(): row[2] for row in values["По дням"].iter_rows(min_row=2, values_only=True)}
    assert days[datetime(2026, 9, 13).date()] == 0
    assert sum(days.values()) == 2
    for row in values["События"].iter_rows(min_row=2, values_only=True):
        if row[8] != "ADMIN":
            assert row[6] in (None, "") and row[7] in (None, "")
    for sheet in values:
        assert all(cell.data_type != "e" for row in sheet for cell in row)


def test_legacy_bot_identity_does_not_invent_a_messenger():
    item = sample_payload()["items"][1]
    item["activities"] = [
        activity(
            1,
            15,
            "STATUS_CHANGED",
            actor_type="UNKNOWN",
            channel="UNKNOWN",
            previous="NEW",
            status="IN_PROGRESS",
        )
    ]
    fact = request_facts(item, {"bot-admin": "Борис"}, {})
    assert fact["first"]["actorChannel"] == "BOT"
    assert fact["first"]["attribution"] == "LEGACY_BOT_ACCOUNT"
    assert fact["hasCreation"] is False
    assert request_facts(item, {}, {})["first"]["actorChannel"] == "UNKNOWN"


def test_full_event_history_is_preserved_and_direct_completion_has_no_work_duration():
    payload = sample_payload()
    item = payload["items"][1]
    item["activities"] = [activity(i, i, "MESSAGE_FROM_USER", actor_type="USER") for i in range(205)]
    item["activities"].append(activity(999, 300, "STATUS_CHANGED", previous="NEW", status="COMPLETED"))
    content = build_statistics_workbook(payload, [], ZoneInfo("Europe/Moscow"))
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    assert wb["События"].max_row == 207
    assert wb["Заявки"]["O2"].value == 300
    assert wb["Заявки"]["V2"].value in (None, "")


def test_empty_export_and_truncated_response():
    payload = sample_payload()
    payload.update(items=[], total=0)
    content = build_statistics_workbook(payload, [], ZoneInfo("Europe/Moscow"))
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    assert wb["Сводка"]["B2"].value == 0
    assert wb["Сводка"]["B8"].value in (None, "")
    payload["total"] = 1
    with pytest.raises(ValueError, match="неполная"):
        build_statistics_workbook(payload, [], ZoneInfo("Europe/Moscow"))
